import streamlit as st
from openai import OpenAI
import pandas as pd
import PyPDF2
from docx import Document
from fpdf import FPDF
from duckduckgo_search import DDGS
from pptx import Presentation
from ics import Calendar, Event
import base64
import io
import zipfile
import json
from datetime import datetime

# --- 1. CONFIGURATION ---
st.set_page_config(page_title="Studio V24: Omni-King", layout="wide", page_icon="👑")

st.markdown("""
    <style>
    .stButton button {height: 3.5em; font-weight: 800; border-radius: 8px; background-color: #000000; color: white; border: 1px solid #333;}
    .brand-sidebar {background-color: #f8f9fa; padding: 15px; border-radius: 10px;}
    div[data-testid="stStatusWidget"] {visibility: visible;}
    </style>
    """, unsafe_allow_html=True)

# --- 2. AUTHENTICATION ---
if "OPENAI_API_KEY" in st.secrets:
    api_key = st.secrets["OPENAI_API_KEY"]
else:
    api_key = st.sidebar.text_input("🔑 API Key", type="password")
    if not api_key: st.stop()

client = OpenAI(api_key=api_key)

if "history" not in st.session_state: st.session_state.history = []

# --- 3. HELPER FUNCTIONS ---
def encode_image(image_file):
    return base64.b64encode(image_file.getvalue()).decode('utf-8')

def create_ics(events_data):
    c = Calendar()
    for e in events_data:
        event = Event()
        event.name = e.get('summary', 'Task')
        try: event.begin = e.get('date', datetime.now().strftime("%Y-%m-%d"))
        except: event.begin = datetime.now().strftime("%Y-%m-%d")
        c.events.add(event)
    return str(c).encode('utf-8')

def create_audio(text):
    try: return client.audio.speech.create(model="tts-1", voice="onyx", input=text[:4096]).content
    except: return None

def create_pptx(slides_content):
    prs = Presentation()
    for slide_data in slides_content:
        layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = slide_data.get('title', 'Untitled')
        slide.placeholders[1].text = slide_data.get('body', '')
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()

def create_pdf(text):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.multi_cell(0, 10, text.encode('latin-1', 'replace').decode('latin-1'))
    return pdf.output(dest="S").encode("latin-1")

def web_search(query):
    try: return "\n".join([f"- {r['title']}: {r['body']}" for r in DDGS().text(query, max_results=2)])
    except: return "Web search offline."

def read_file(uploaded_file):
    if not uploaded_file: return ""
    try:
        if uploaded_file.name.endswith('.pdf'): return "".join([p.extract_text() for p in PyPDF2.PdfReader(uploaded_file).pages])
        elif uploaded_file.name.endswith('.docx'): return "\n".join([p.text for p in Document(uploaded_file).paragraphs])
        elif uploaded_file.name.endswith('.txt'): return uploaded_file.read().decode("utf-8")
    except: return ""

# --- 4. THE OMNI ENGINE ---

def run_omni_swarm(mission, context, image_data, voice_transcript, req_outputs, depth):
    results = {}
    
    with st.status(f"👑 Executing {depth} Mission...", expanded=True) as status:
        
        # A. SENSORY INPUT (Web + Vision + Voice)
        status.write("👁️ Agent 1: Sensory Analysis (Web/Vision/Voice)...")
        intel = web_search(f"{mission} trends data")
        
        visual_analysis = ""
        if image_data:
            b64_img = encode_image(image_data)
            vis_resp = client.chat.completions.create(
                model="gpt-4o",
                messages=[{"role": "user", "content": [{"type": "text", "text": "Analyze for relevance."}, {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{b64_img}"}}]}]
            )
            visual_analysis = f"VISUAL INTEL: {vis_resp.choices[0].message.content}"
        
        # B. CORE STRATEGY
        status.write(f"🧠 Agent 2: Strategic Synthesis...")
        depth_prompts = {
            "Draft (Speed)": "Concise. Bullet points.",
            "Standard (Balanced)": "Standard professional depth.",
            "Deep Dive (Comprehensive)": "Extensive analysis, risks, and competitor breakdown."
        }
        
        full_context = f"MISSION: {mission}\nVOICE NOTE: {voice_transcript}\nINTEL: {intel}\n{visual_analysis}\nCONTEXT: {context}\nDEPTH: {depth_prompts[depth]}"
        
        # Recursive Drafting
        attempts = 2 if depth == "Deep Dive (Comprehensive)" else 1
        current_draft = ""
        for i in range(attempts):
            draft_resp = client.chat.completions.create(
                model="gpt-4o",
                messages=[{"role": "system", "content": f"Execute Mission. Output Markdown."}, {"role": "user", "content": f"Refine: {current_draft}" if current_draft else f"Context: {full_context}"}]
            )
            current_draft = draft_resp.choices[0].message.content
        
        results['Master_Strategy.md'] = current_draft
        
        # C. ARTIFACTS
        if req_outputs:
            status.write("📦 Agent 3: Manufacturing Assets...")
            
            # 1. Sim Log
            if "Simulation Log" in req_outputs:
                sim_resp = client.chat.completions.create(
                    model="gpt-4o", messages=[{"role": "system", "content": f"Simulate focus group reaction to: {current_draft[:1000]}"}]
                )
                results['Simulation_Log.md'] = sim_resp.choices[0].message.content

            # 2. PPTX
            if "Presentation (.pptx)" in req_outputs:
                try:
                    slide_resp = client.chat.completions.create(
                        model="gpt-4o", response_format={ "type": "json_object" },
                        messages=[{"role": "system", "content": f"Convert to 5 slides JSON: {current_draft[:2000]}"}]
                    )
                    results['Presentation.pptx'] = create_pptx(json.loads(slide_resp.choices[0].message.content).get('slides', []))
                except: pass
            
            # 3. Calendar
            if "Calendar (.ics)" in req_outputs:
                try:
                    cal_resp = client.chat.completions.create(
                        model="gpt-4o", response_format={ "type": "json_object" },
                        messages=[{"role": "system", "content": f"Extract dates to JSON: {current_draft[:2000]}"}]
                    )
                    results['Schedule.ics'] = create_ics(json.loads(cal_resp.choices[0].message.content).get('events', []))
                except: pass

            # 4. Audio
            if "Audio Brief (.mp3)" in req_outputs:
                results['Audio_Brief.mp3'] = create_audio(f"Briefing. {current_draft[:500]}")

            # 5. RESTORED: Concept Art
            if "Concept Art (.png)" in req_outputs:
                try:
                    art_p = client.chat.completions.create(
                        model="gpt-4o", messages=[{"role": "system", "content": f"Create DALL-E 3 prompt for: {mission}"}]
                    ).choices[0].message.content
                    results['Concept_Art_URL'] = client.images.generate(model="dall-e-3", prompt=art_p, size="1024x1024").data[0].url
                except: pass

            # 6. RESTORED: Data Tables
            if "Data Tables (.csv)" in req_outputs:
                csv_resp = client.chat.completions.create(
                    model="gpt-4o", messages=[{"role": "system", "content": f"Extract key data/tasks from this strategy into a CSV format. Output ONLY the CSV content.\n\nSTRATEGY: {current_draft[:2000]}"}]
                )
                results['Project_Data.csv'] = csv_resp.choices[0].message.content

        status.update(label="✅ Omni-Cycle Complete", state="complete")
        
    return results

# --- 5. INTERFACE ---
with st.sidebar:
    st.title("🧬 Brand DNA")
    brand_name = st.text_input("Org Name", placeholder="Apex Media")
    north_star = st.text_area("North Star Goal", placeholder="Growth...")

st.title("👑 Studio V24: The Omni-King")

# UNIFIED INPUT (Restored Voice)
tab_text, tab_vis, tab_voice = st.tabs(["📝 Mission Control", "👁️ Vision Uplink", "🎙️ Voice Command"])
with tab_text:
    mission_input = st.text_area("Objective", placeholder="e.g. Launch Q3 Strategy...", height=100)
    uploaded_file = st.file_uploader("Intel Files", type=["pdf", "docx", "txt"])
with tab_vis:
    camera_img = st.camera_input("Scan Target")
    uploaded_img = st.file_uploader("Upload Visual", type=["jpg", "png"])
with tab_voice:
    voice_cmd = st.audio_input("Record Instruction")

st.divider()

col1, col2 = st.columns([1, 1])
with col1:
    st.subheader("🎚️ Depth Control")
    depth_setting = st.select_slider("Analysis Depth", options=["Draft (Speed)", "Standard (Balanced)", "Deep Dive (Comprehensive)"], value="Standard (Balanced)")

with col2:
    st.subheader("📦 Output Selection")
    selected_outputs = st.multiselect(
        "Select Artifacts:",
        ["Simulation Log", "Presentation (.pptx)", "Calendar (.ics)", "Audio Brief (.mp3)", "Concept Art (.png)", "Data Tables (.csv)"],
        default=["Presentation (.pptx)", "Calendar (.ics)", "Concept Art (.png)"]
    )

if st.button("🚀 IGNITE ENGINE", type="primary", use_container_width=True):
    # Inputs
    active_img = camera_img if camera_img else uploaded_img
    file_txt = read_file(uploaded_file)
    
    # Transcribe Voice
    voice_txt = ""
    if voice_cmd:
        voice_txt = client.audio.transcriptions.create(model="whisper-1", file=voice_cmd).text
        st.info(f"🎙️ Heard: {voice_txt}")
    
    context = f"Brand: {brand_name}. Goal: {north_star}. File Data: {file_txt}"
    
    # Execute
    if mission_input or voice_txt or active_img:
        outputs = run_omni_swarm(mission_input, context, active_img, voice_txt, selected_outputs, depth_setting)
        st.session_state.history.insert(0, outputs)
        st.rerun()
    else:
        st.warning("⚠️ Input Required (Text, Voice, or Image)")

# --- 6. OUTPUTS ---
if st.session_state.history:
    st.divider()
    latest = st.session_state.history[0]
    
    # DYNAMIC DASHBOARD
    if 'Concept_Art_URL' in latest:
        st.image(latest['Concept_Art_URL'], caption="Visual Concept", width=400)
        
    tabs = ["📄 Strategy"]
    if "Simulation_Log.md" in latest: tabs.append("👥 Sim")
    if "Audio_Brief.mp3" in latest: tabs.append("🎧 Audio")
    tabs.append("📦 Download")
    
    active_tabs = st.tabs(tabs)
    
    with active_tabs[0]: st.markdown(latest.get('Master_Strategy.md', ''))
    if "👥 Sim" in tabs:
        with active_tabs[tabs.index("👥 Sim")]: st.info(latest['Simulation_Log.md'])
    if "🎧 Audio" in tabs:
        with active_tabs[tabs.index("🎧 Audio")]: st.audio(latest['Audio_Brief.mp3'])
            
    with active_tabs[-1]:
        st.success("Assets Ready")
        zip_buffer = io.BytesIO()
        with zipfile.ZipFile(zip_buffer, "a", zipfile.ZIP_DEFLATED, False) as zip_file:
            for key, val in latest.items():
                if key.endswith(('.pptx', '.mp3', '.ics')): zip_file.writestr(key, val)
                elif key.endswith('.md') or key.endswith('.csv'): 
                    zip_file.writestr(key, val)
                    if key.endswith('.md'):
                        try: zip_file.writestr(key.replace('.md', '.pdf'), create_pdf(val))
                        except: pass
        
        st.download_button("📦 Download Bundle (.ZIP)", zip_buffer.getvalue(), "Omni_Assets.zip", "application/zip", type="primary", use_container_width=True)
